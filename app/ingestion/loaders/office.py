# pyrefly: ignore [missing-import]
# pyrefly: ignore [missing-import]
# pyrefly: ignore [missing-import]
import os
# pyrefly: ignore [missing-import]
import logfire
# pyrefly: ignore [missing-import]
from docx import Document
# pyrefly: ignore [missing-import]
from pptx import Presentation


def _parse_docx(file_path: str) -> str:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables, which python-docx skips by default
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def _parse_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    slide_texts = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        texts.append(text)
        if texts:
            slide_texts.append(f"[Slide {slide_num}]\n" + "\n".join(texts))

    return "\n\n".join(slide_texts)


def parse_office(file_path: str) -> str:
    """
    Parses Word (.docx) and PowerPoint (.pptx) files and extracts text content.
    """
    with logfire.span("Office Parsing", filename=file_path):
        try:
            ext = os.path.splitext(file_path)[1].lower()

            if ext == ".docx":
                return _parse_docx(file_path)
            elif ext == ".pptx":
                return _parse_pptx(file_path)
            else:
                raise ValueError(f"Unsupported office file type: {ext}")

        except Exception as e:
            logfire.error(f"Office Parse failed: {e}")
            raise e